"""
Export Utility for Polished Leads
Exports polished lead intelligence to Excel (.xlsx) and CSV with proper formatting.
"""

from pathlib import Path
from typing import List, Dict, Any
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXPORT_COLUMNS = [
    ("business_name", "Business Name", 30),
    ("category", "Category", 22),
    ("owner_name", "Owner / Decision Maker", 25),
    ("owner_role", "Role / Title", 18),
    ("confidence_score", "Confidence (%)", 14),
    ("phone", "Direct Phone", 18),
    ("email", "Verified Email", 28),
    ("instagram", "Instagram Profile", 32),
    ("facebook", "Facebook Page", 32),
    ("linkedin", "LinkedIn Profile", 32),
    ("city", "City", 18),
    ("state", "State", 8),
    ("address", "Full Address", 38),
    ("rating", "Rating", 10),
    ("reviews", "Reviews", 10),
    ("has_website", "Has Website?", 14),
    ("pitch_hook", "High-Ticket Pitch Hook", 55),
    ("sources_scanned", "Sources Scanned", 35),
    ("scraped_at", "Scraped Date", 20)
]

def export_leads(leads: List[Dict[str, Any]], base_path: Path):
    """
    Saves leads to both .xlsx and .csv at base_path (without extension).
    Applies professional styling to the Excel sheet.
    """
    if not leads:
        print("[EXPORT] No leads provided for export.")
        return

    # Normalize records into standard table structure
    rows = []
    for lead in leads:
        row = {}
        for key, col_name, _ in EXPORT_COLUMNS:
            row[col_name] = lead.get(key, "")
        rows.append(row)

    df = pd.DataFrame(rows)

    # 1. Export CSV
    csv_file = base_path.with_suffix(".csv")
    df.to_csv(csv_file, index=False, encoding="utf-8")
    print(f"[EXPORT] Saved {len(leads)} leads to CSV: {csv_file}")

    # 2. Export Excel with OpenPyXL styling
    xlsx_file = base_path.with_suffix(".xlsx")
    with pd.ExcelWriter(xlsx_file, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="High-Ticket Leads", index=False)
        ws = writer.sheets["High-Ticket Leads"]

        # Header styling
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        border_style = Side(border_style="thin", color="D9D9D9")
        thin_border = Border(left=border_style, right=border_style, top=border_style, bottom=border_style)

        for col_num, (_, _, width) in enumerate(EXPORT_COLUMNS, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            col_letter = get_column_letter(col_num)
            ws.column_dimensions[col_letter].width = max(width, 12)

        # Body alignment and freeze top row
        ws.views.sheetView[0].tabSelected = True
        ws.freeze_panes = "A2"

        for row in ws.iter_rows(min_row=2, max_row=len(leads) + 1, min_col=1, max_col=len(EXPORT_COLUMNS)):
            for cell in row:
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)
                cell.alignment = Alignment(vertical="center")

    print(f"[EXPORT] Saved styled Excel workbook: {xlsx_file}")
